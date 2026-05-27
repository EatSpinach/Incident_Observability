"""
Incident Observability Web Application
Flask-based web interface for evaluating incident worknotes and closing comments
"""

from flask import Flask, render_template, request, jsonify, send_file, session
from flask_cors import CORS
import os
import json
from datetime import datetime
from collections import Counter
import openpyxl
import xlrd
from werkzeug.utils import secure_filename
import io
import traceback

app = Flask(__name__)
CORS(app)
app.secret_key = os.environ.get('SECRET_KEY', 'dev-secret-key-change-in-production')
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max file size
app.config['UPLOAD_FOLDER'] = 'uploads'

# Ensure upload folder exists
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Load configuration
def load_config():
    """Load configuration from config.json"""
    try:
        with open('config.json', 'r') as f:
            return json.load(f)
    except Exception as e:
        print(f"Error loading config: {e}")
        return get_default_config()

def get_default_config():
    """Return default configuration"""
    return {
        "worknotes": {
            "required_keywords": ["detection", "analysis", "impact", "log", "monitoring"]
        },
        "closing_comments": {
            "required_keywords": ["resolution", "verification", "rootcause"]
        },
        "scoring": {
            "keywords_weight": 1.0
        },
        "evaluation_criteria": {
            "good_threshold": 74,
            "average_threshold": 50,
            "poor_threshold": 0
        },
        "phase_keywords": {
            "detection": ["detection", "detected", "identified", "discovered", "found", "alarm", "alert", "detect", "identify", "discover"],
            "analysis": ["analysis", "analyzed", "investigated", "examined", "reviewed", "analyze", "investigate", "examine", "review", "troubleshoot", "check"],
            "impact": ["impact", "affected", "severity", "priority", "affect", "severe", "urgent", "urgency", "critical", "complex"],
            "log": ["log", "logs", "logged", "logging"],
            "rootcause": ["root cause", "rootcause", "root-cause", "cause identified", "reason", "rca", "cause", "problem", "issue"],
            "resolution": ["resolution", "resolved", "fixed", "corrected", "remediated", "fix", "solution", "correction"],
            "verification": ["verification", "verified", "tested", "confirmed", "validated", "checked"],
            "monitoring": ["monitoring", "monitor", "monitored", "tracking", "observing", "watching"]
        }
    }

config = load_config()

def safe_cell_text(value):
    """Convert Excel cell values to safe text"""
    return "" if value is None else str(value)

def read_excel_data(file_path):
    """Read Excel data"""
    extension = os.path.splitext(file_path)[1].lower()
    
    if extension == '.xlsx':
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        sheet = workbook.active
        if sheet is None:
            return []
        rows = list(sheet.iter_rows(values_only=True))
    elif extension == '.xls':
        workbook = xlrd.open_workbook(file_path)
        sheet = workbook.sheet_by_index(0)
        rows = [sheet.row_values(i) for i in range(sheet.nrows)]
    else:
        raise ValueError("Unsupported file format. Please select .xlsx or .xls")
    
    if not rows:
        return []
    
    headers = [str(cell) if cell is not None else f"Column{idx+1}" for idx, cell in enumerate(rows[0])]
    data = []
    for row in rows[1:]:
        row_dict = {}
        for idx, header in enumerate(headers):
            row_dict[header] = row[idx] if idx < len(row) else None
        data.append(row_dict)
    return data

def detect_columns(df):
    """Detect worknotes and closing comments columns"""
    if not df:
        return None
    
    worknotes_col = None
    closing_col = None
    incident_col = None
    app_group_col = None
    resolved_by_col = None
    
    for col in df[0].keys():
        if any(keyword in col.lower() for keyword in ['worknote', 'work note', 'notes', 'comments']):
            if 'closing' not in col.lower() and 'close' not in col.lower():
                worknotes_col = col
                break
    
    for col in df[0].keys():
        if any(keyword in col.lower() for keyword in ['closing', 'close comment', 'resolution']):
            closing_col = col
            break
    
    for col in df[0].keys():
        if any(keyword in col.lower() for keyword in ['incident', 'id', 'number', 'ticket']):
            incident_col = col
            break
    
    for col in df[0].keys():
        if any(keyword in col.lower() for keyword in ['application', 'app', 'group', 'assignment', 'team', 'category']):
            app_group_col = col
            break
    
    for col in df[0].keys():
        if 'resolved by' in col.lower():
            resolved_by_col = col
            break
    
    if not worknotes_col and not closing_col:
        return None
    
    return {
        'incident': incident_col or list(df[0].keys())[0],
        'worknotes': worknotes_col,
        'closing': closing_col,
        'app_group': app_group_col,
        'resolved_by': resolved_by_col
    }

def evaluate_text(text, criteria, config):
    """Evaluate text based on criteria"""
    score = 0
    issues = []
    text_lower = text.lower()
    
    phase_keywords = config.get('phase_keywords', {})
    
    found_required = 0
    missing_keywords = []
    
    for kw in criteria['required_keywords']:
        kw_normalized = kw.lower().replace(' ', '_')
        synonyms = phase_keywords.get(kw_normalized, [kw.lower()])
        
        if any(syn.lower() in text_lower for syn in synonyms):
            found_required += 1
        else:
            missing_keywords.append(kw)
    
    keyword_score = (found_required / len(criteria['required_keywords'])) * 100
    
    if found_required == 0:
        issues.append("No required keywords found")
    elif found_required < len(criteria['required_keywords']) / 2:
        issues.append(f"Only {found_required}/{len(criteria['required_keywords'])} required keywords found")
    
    score = keyword_score * config['scoring']['keywords_weight']
    
    return score, issues

def evaluate_row(row, cols, config):
    """Evaluate a single incident row"""
    incident_id = row.get(cols['incident'], 'N/A') if cols['incident'] else 'N/A'
    input_number = safe_cell_text(row.get('Number', ''))
    
    app_group = 'Ungrouped'
    if cols.get('app_group'):
        app_group = safe_cell_text(row.get(cols['app_group'], 'Ungrouped'))
        app_group = app_group.strip() if app_group.strip() else 'Ungrouped'
    
    worknotes_score = 0
    worknotes_issues = []
    if cols['worknotes']:
        worknotes = safe_cell_text(row.get(cols['worknotes']))
        worknotes_score, worknotes_issues = evaluate_text(worknotes, config['worknotes'], config)
    
    closing_score = 0
    closing_issues = []
    if cols['closing']:
        closing = safe_cell_text(row.get(cols['closing']))
        closing_score, closing_issues = evaluate_text(closing, config['closing_comments'], config)
    
    overall_score = (worknotes_score + closing_score) / 2 if cols['worknotes'] and cols['closing'] else max(worknotes_score, closing_score)
    
    if overall_score > 74:
        overall_rating = "Good"
    elif overall_score >= 50 and overall_score < 75:
        overall_rating = "Average"
    else:
        overall_rating = "Poor"
    
    if worknotes_score > 74:
        worknotes_rating = "Good"
    elif worknotes_score >= 50 and worknotes_score < 75:
        worknotes_rating = "Average"
    else:
        worknotes_rating = "Poor"
        
    if closing_score > 74:
        closing_rating = "Good"
    elif closing_score >= 50 and closing_score < 75:
        closing_rating = "Average"
    else:
        closing_rating = "Poor"
    
    return {
        'Number': input_number,
        'Group': app_group,
        'Resolved by': safe_cell_text(row.get('Resolved by', '')),
        'Resolved': safe_cell_text(row.get('Resolved', '')),
        'Worknotes Score': round(worknotes_score, 2),
        'Worknotes Rating': worknotes_rating,
        'Closing Comments Score': round(closing_score, 2),
        'Closing Comments Rating': closing_rating,
        'Total Score': round(overall_score, 2),
        'Total Rating': overall_rating,
        'Worknotes Issues': '; '.join(worknotes_issues) if worknotes_issues else 'None',
        'Closing Comments Issues': '; '.join(closing_issues) if closing_issues else 'None'
    }

@app.route('/')
def index():
    """Home page"""
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    """Handle file upload and evaluation"""
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'Invalid file format. Please upload .xlsx or .xls file'}), 400
    
    try:
        filename = secure_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(filepath)
        
        # Read and process data
        df = read_excel_data(filepath)
        cols = detect_columns(df)
        
        if not cols:
            os.remove(filepath)
            return jsonify({'error': 'Could not detect required columns (Worknotes/Closing Comments)'}), 400
        
        # Evaluate incidents
        results = []
        for row in df:
            result = evaluate_row(row, cols, config)
            results.append(result)
        
        # Store in session
        session['results'] = results
        session['filename'] = filename
        
        # Calculate statistics
        stats = calculate_statistics(results)
        
        # Clean up uploaded file
        os.remove(filepath)
        
        return jsonify({
            'success': True,
            'results': results,
            'statistics': stats,
            'total_records': len(results)
        })
    
    except Exception as e:
        print(f"Error in upload_file: {str(e)}")
        print(traceback.format_exc())
        return jsonify({'error': f'Error processing file: {str(e)}'}), 500

def calculate_statistics(results):
    """Calculate statistics from evaluation results"""
    if not results:
        return {}
    
    worknotes_scores = [r['Worknotes Score'] for r in results]
    closing_scores = [r['Closing Comments Score'] for r in results]
    total_scores = [r['Total Score'] for r in results]
    
    worknotes_ratings = Counter(r['Worknotes Rating'] for r in results)
    closing_ratings = Counter(r['Closing Comments Rating'] for r in results)
    total_ratings = Counter(r['Total Rating'] for r in results)
    
    return {
        'total_incidents': len(results),
        'worknotes_avg': round(sum(worknotes_scores) / len(worknotes_scores), 2),
        'closing_avg': round(sum(closing_scores) / len(closing_scores), 2),
        'overall_avg': round(sum(total_scores) / len(total_scores), 2),
        'worknotes_good': worknotes_ratings.get('Good', 0),
        'worknotes_average': worknotes_ratings.get('Average', 0),
        'worknotes_poor': worknotes_ratings.get('Poor', 0),
        'closing_good': closing_ratings.get('Good', 0),
        'closing_average': closing_ratings.get('Average', 0),
        'closing_poor': closing_ratings.get('Poor', 0),
        'total_good': total_ratings.get('Good', 0),
        'total_average': total_ratings.get('Average', 0),
        'total_poor': total_ratings.get('Poor', 0)
    }

@app.route('/export')
def export_results():
    """Export results to Excel"""
    results = session.get('results', [])
    if not results:
        return jsonify({'error': 'No results to export'}), 400
    
    try:
        output = io.BytesIO()
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Evaluation Results"
        
        # Headers
        headers = ['S.No.', 'Number', 'Group', 'Resolved by', 'Resolved', 
                   'Worknotes Score', 'Worknotes Rating', 
                   'Closing Comments Score', 'Closing Comments Rating',
                   'Total Score', 'Total Rating', 
                   'Worknotes Issues', 'Closing Comments Issues']
        sheet.append(headers)
        
        # Data
        for idx, result in enumerate(results, start=1):
            row = [
                idx,
                result.get('Number', ''),
                result.get('Group', ''),
                result.get('Resolved by', ''),
                result.get('Resolved', ''),
                result.get('Worknotes Score', 0),
                result.get('Worknotes Rating', ''),
                result.get('Closing Comments Score', 0),
                result.get('Closing Comments Rating', ''),
                result.get('Total Score', 0),
                result.get('Total Rating', ''),
                result.get('Worknotes Issues', ''),
                result.get('Closing Comments Issues', '')
            ]
            sheet.append(row)
        
        workbook.save(output)
        output.seek(0)
        
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'evaluation_results_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        )
    
    except Exception as e:
        return jsonify({'error': f'Error exporting results: {str(e)}'}), 500

@app.route('/config')
def get_config():
    """Get current configuration"""
    return jsonify(config)

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)

# Made with Bob
